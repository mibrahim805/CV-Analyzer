from io import BytesIO

from docx import Document
from pypdf import PdfWriter

from app.services.document_extraction import extract_document


def test_extracts_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("Backend engineer with five years of experience.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "PostgreSQL"
    content = BytesIO()
    document.save(content)

    result = extract_document(
        content.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert result.method == "docx_text"
    assert result.needs_ocr is False
    assert "Backend engineer" in result.text
    assert "Python | PostgreSQL" in result.text


def test_empty_pdf_is_marked_for_ocr() -> None:
    content = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(content)

    result = extract_document(content.getvalue(), "application/pdf")

    assert result.method == "ocr_required"
    assert result.needs_ocr is True
    assert result.page_count == 1


def test_rejects_unsupported_media_type() -> None:
    try:
        extract_document(b"content", "text/plain")
    except ValueError as exc:
        assert "Only PDF and DOCX" in str(exc)
    else:
        raise AssertionError("unsupported media type was accepted")
