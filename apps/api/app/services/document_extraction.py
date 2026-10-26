"""Deterministic text extraction for approved CV document formats.

OCR deliberately remains a separate capability. A document with no embedded text
is marked as requiring OCR rather than being guessed or silently discarded.
"""

import re
from dataclasses import dataclass
from io import BytesIO


class ExtractionError(ValueError):
    def __init__(self, code: str, detail: str) -> None:
        super().__init__(detail)
        self.code = code
        self.detail = detail


@dataclass(frozen=True)
class ExtractionResult:
    method: str
    text: str
    page_count: int | None
    needs_ocr: bool = False


def extract_document(content: bytes, media_type: str) -> ExtractionResult:
    if media_type == "application/pdf":
        return _extract_pdf(content)
    if media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _extract_docx(content)
    raise ExtractionError("unsupported_media_type", "Only PDF and DOCX documents can be extracted.")


def _extract_pdf(content: bytes) -> ExtractionResult:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError("extractor_unavailable", "PDF extraction is not available on this worker.") from exc
    try:
        reader = PdfReader(BytesIO(content), strict=False)
    except Exception as exc:  # pypdf exposes several parser-specific exception types.
        raise ExtractionError("malformed_document", "The PDF could not be parsed safely.") from exc
    if reader.is_encrypted:
        try:
            if reader.decrypt("") == 0:
                raise ExtractionError("encrypted_document", "This PDF is encrypted and needs to be unlocked first.")
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError(
                "encrypted_document", "This PDF is encrypted and needs to be unlocked first."
            ) from exc

    pages: list[str] = []
    try:
        for page in reader.pages:
            pages.append(page.extract_text() or "")
    except Exception as exc:
        raise ExtractionError("malformed_document", "The PDF text could not be read safely.") from exc
    text = _normalize_text("\n\n".join(pages))
    if not text:
        return ExtractionResult(method="ocr_required", text="", page_count=len(reader.pages), needs_ocr=True)
    return ExtractionResult(method="pdf_text", text=text, page_count=len(reader.pages))


def _extract_docx(content: bytes) -> ExtractionResult:
    try:
        from docx import Document
    except ImportError as exc:
        raise ExtractionError("extractor_unavailable", "DOCX extraction is not available on this worker.") from exc
    try:
        document = Document(BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text for cell in row.cells))
    except Exception as exc:
        raise ExtractionError("malformed_document", "The DOCX could not be parsed safely.") from exc
    text = _normalize_text("\n".join(blocks))
    if not text:
        return ExtractionResult(method="ocr_required", text="", page_count=None, needs_ocr=True)
    return ExtractionResult(method="docx_text", text=text, page_count=None)


def _normalize_text(value: str) -> str:
    value = value.replace("\x00", "")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()
